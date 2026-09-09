from __future__ import annotations

import asyncio
import io
import struct
import threading
import zipfile
from collections.abc import Callable

import pytest
from docx import Document
from openpyxl import Workbook
from openpyxl.worksheet.formula import ArrayFormula
from pptx import Presentation
from pptx.util import Inches

from cairn.ingestion.base import ParseContext, ParsedDocument
from cairn.ingestion.errors import ParseCorruptFile, ParseEmptyContent
from cairn.ingestion.office import CsvParser, DocxParser, PptxParser, SheetParser, XlsxParser
from cairn.ingestion.office_package import inspect_ooxml_package


def _bytes_from(save: Callable[[io.BytesIO], object]) -> bytes:
    stream = io.BytesIO()
    save(stream)
    return stream.getvalue()


def _docx_fixture() -> bytes:
    document = Document()
    document.add_heading("Overview", level=1)
    document.add_paragraph("Opening paragraph.")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Name"
    table.cell(0, 1).text = "Value"
    table.cell(1, 0).text = "alpha"
    table.cell(1, 1).text = "1 | 2"
    document.add_heading("Details", level=2)
    document.add_paragraph("Nested item", style="List Bullet")
    return _bytes_from(document.save)


def _pptx_fixture() -> bytes:
    presentation = Presentation()
    first = presentation.slides.add_slide(presentation.slide_layouts[1])
    first.shapes.title.text = "Summary"
    first.placeholders[1].text = "First point\nSecond point"
    second = presentation.slides.add_slide(presentation.slide_layouts[5])
    second.shapes.title.text = "Metrics"
    frame = second.shapes.add_table(
        2,
        2,
        Inches(1),
        Inches(2),
        Inches(4),
        Inches(1),
    )
    frame.table.cell(0, 0).text = "Metric"
    frame.table.cell(0, 1).text = "Value"
    frame.table.cell(1, 0).text = "quality"
    frame.table.cell(1, 1).text = "high"
    return _bytes_from(presentation.save)


def _xlsx_fixture() -> bytes:
    workbook = Workbook()
    active = workbook.active
    assert active is not None
    active.title = "Data"
    active.append(["Name", "Total"])
    active.append(["alpha", "=SUM(2,3)"])
    hidden = workbook.create_sheet("Hidden")
    hidden.sheet_state = "hidden"
    hidden.append(["Secret", "Value"])
    hidden.append(["token", 7])
    return _bytes_from(workbook.save)


def _zip_with(name: str, payload: bytes) -> bytes:
    stream = io.BytesIO()
    with zipfile.ZipFile(stream, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("[Content_Types].xml", "<Types/>")
        archive.writestr(name, payload)
    return stream.getvalue()


def _replace_member(package: bytes, name: str, payload: bytes) -> bytes:
    output = io.BytesIO()
    with (
        zipfile.ZipFile(io.BytesIO(package)) as source,
        zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as target,
    ):
        for member in source.infolist():
            if member.filename != name:
                target.writestr(member, source.read(member))
        target.writestr(name, payload)
    return output.getvalue()


async def test_docx_preserves_document_order_headings_tables_and_source() -> None:
    parsed = await DocxParser().parse(
        _docx_fixture(),
        ParseContext(source_url="https://example.test/report.docx", language="en"),
    )

    assert [block.kind for block in parsed.blocks] == [
        "heading",
        "paragraph",
        "table",
        "heading",
        "list",
    ]
    assert parsed.blocks[1].heading_path == ("Overview",)
    assert parsed.blocks[2].text == "| Name | Value |\n| --- | --- |\n| alpha | 1 \\| 2 |"
    assert parsed.blocks[4].heading_path == ("Overview", "Details")
    assert all(block.page is None for block in parsed.blocks)
    assert all(block.source_url == "https://example.test/report.docx" for block in parsed.blocks)
    assert parsed.pages == 1
    assert parsed.language == "en"


async def test_pptx_emits_one_atomic_block_per_slide_with_table_provenance() -> None:
    parsed = await PptxParser().parse(_pptx_fixture(), ParseContext())

    assert len(parsed.blocks) == 2
    assert parsed.pages == 2
    assert parsed.blocks[0].page == 1
    assert parsed.blocks[0].heading_path == ("Summary",)
    assert "First point" in parsed.blocks[0].text
    assert parsed.blocks[1].kind == "table"
    assert parsed.blocks[1].page == 2
    assert parsed.blocks[1].heading_path == ("Metrics",)
    assert "| Metric | Value |" in parsed.blocks[1].text
    assert parsed.metadata["slides"] == [
        {"number": 1, "title": "Summary"},
        {"number": 2, "title": "Metrics"},
    ]


async def test_xlsx_preserves_headers_formulas_sheet_order_and_visibility() -> None:
    parsed = await SheetParser().parse(_xlsx_fixture(), ParseContext())

    assert len(parsed.blocks) == 2
    assert parsed.pages == 2
    assert parsed.blocks[0].kind == "table"
    assert parsed.blocks[0].heading_path == ("Data",)
    assert parsed.blocks[0].page == 1
    assert "| Name | Total |" in parsed.blocks[0].text
    assert "| alpha | =SUM(2,3) |" in parsed.blocks[0].text
    assert parsed.blocks[1].heading_path == ("Hidden",)
    assert parsed.blocks[1].page == 2
    assert parsed.metadata["sheets"] == [
        {"name": "Data", "state": "visible"},
        {"name": "Hidden", "state": "hidden"},
    ]
    assert parsed.metadata["formula_mode"] == "source"


async def test_csv_uses_explicit_encoding_and_rectangularizes_ragged_rows() -> None:
    data = "名称;值\r\n甲;一|二\r\n乙\r\n".encode("gb18030")
    parsed = await SheetParser().parse(
        data,
        ParseContext(encoding="gb18030", language="zh"),
    )

    assert parsed.markdown == ("| 名称 | 值 |\n| --- | --- |\n| 甲 | 一\\|二 |\n| 乙 |  |")
    assert parsed.blocks[0].page == 1
    assert parsed.metadata["delimiter"] == ";"
    assert parsed.language == "zh"


@pytest.mark.parametrize("encoding", ["utf-8", "not-an-encoding"])
async def test_csv_rejects_invalid_explicit_decoding(encoding: str) -> None:
    with pytest.raises(ParseCorruptFile):
        await SheetParser().parse(b"\xff\xfe\x00", ParseContext(encoding=encoding))


@pytest.mark.parametrize(
    "parser,data",
    [
        (DocxParser(), b"not a package"),
        (PptxParser(), b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1encrypted"),
        (SheetParser(), b"PK broken package"),
    ],
    ids=["broken-docx", "encrypted-pptx", "broken-sheet"],
)
async def test_broken_and_encrypted_packages_fail_explicitly(
    parser: DocxParser | PptxParser | SheetParser,
    data: bytes,
) -> None:
    with pytest.raises(ParseCorruptFile) as caught:
        await parser.parse(data, ParseContext())
    assert caught.value.code == "PARSE_CORRUPT_FILE"
    assert caught.value.retryable is False


@pytest.mark.parametrize(
    "name,payload",
    [
        ("../word/document.xml", b"<document/>"),
        ("word/document.xml", b"<!DOCTYPE x [<!ENTITY y 'boom'>]><document>&y;</document>"),
        (
            "word/_rels/document.xml.rels",
            b'<Relationships><Relationship TargetMode="External" '
            b'Target="https://example.test/payload"/></Relationships>',
        ),
        ("word/document.xml", b"A" * 200_000),
    ],
    ids=["traversal", "entity", "external-link", "bomb"],
)
async def test_ooxml_preflight_rejects_unsafe_archives(name: str, payload: bytes) -> None:
    with pytest.raises(ParseCorruptFile):
        await DocxParser().parse(_zip_with(name, payload), ParseContext())


@pytest.mark.parametrize(
    "parser,data",
    [
        (DocxParser(), _bytes_from(Document().save)),
        (PptxParser(), _bytes_from(Presentation().save)),
        (SheetParser(), b"\r\n"),
    ],
    ids=["docx", "pptx", "csv"],
)
async def test_empty_office_documents_fail_consistently(
    parser: DocxParser | PptxParser | SheetParser,
    data: bytes,
) -> None:
    with pytest.raises(ParseEmptyContent):
        await parser.parse(data, ParseContext())


async def test_office_parsing_runs_outside_the_event_loop_thread(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    parser = DocxParser()
    event_loop_thread = threading.get_ident()
    parser_threads: list[int] = []
    original = parser._parse

    def recording_parse(data: bytes, ctx: ParseContext) -> ParsedDocument:
        parser_threads.append(threading.get_ident())
        return original(data, ctx)

    monkeypatch.setattr(parser, "_parse", recording_parse)
    parsed, marker = await asyncio.gather(
        parser.parse(_docx_fixture(), ParseContext()),
        asyncio.sleep(0, result="responsive"),
    )

    assert parsed.blocks
    assert marker == "responsive"
    assert parser_threads and parser_threads[0] != event_loop_thread


@pytest.mark.parametrize("name", ["\\word\\evil.xml", "/evil.xml", "C:/evil.xml", "../evil.xml"])
def test_zip_paths_are_rejected_before_library_loading(name: str) -> None:
    with pytest.raises(ParseCorruptFile, match="path"):
        inspect_ooxml_package(_replace_member(_docx_fixture(), name, b"<root/>"))


def test_utf16_dtd_is_rejected_by_defused_parser() -> None:
    payload = '<?xml version="1.0" encoding="utf-16"?><!DOCTYPE x><x/>'.encode("utf-16")
    with pytest.raises(ParseCorruptFile, match="XML"):
        inspect_ooxml_package(_replace_member(_docx_fixture(), "customXml/test.xml", payload))


@pytest.mark.parametrize("target", ["../../../../secret", "file:///secret", "//server/share"])
def test_relationship_cannot_escape_the_package(target: str) -> None:
    payload = (
        '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
        f'<Relationship Id="rId1" Type="test" Target="{target}"/></Relationships>'
    ).encode()
    with pytest.raises(ParseCorruptFile):
        inspect_ooxml_package(_replace_member(_docx_fixture(), "word/_rels/test.xml.rels", payload))


def test_declared_xml_content_type_is_inspected_even_without_xml_suffix() -> None:
    package = _docx_fixture()
    with zipfile.ZipFile(io.BytesIO(package)) as archive:
        manifest = archive.read("[Content_Types].xml").replace(
            b"</Types>",
            b'<Override PartName="/custom.part" ContentType="application/xml"/></Types>',
        )
    package = _replace_member(package, "[Content_Types].xml", manifest)
    package = _replace_member(package, "custom.part", b"<!DOCTYPE x><x/>")
    with pytest.raises(ParseCorruptFile, match="XML"):
        inspect_ooxml_package(package)


async def test_sparse_sheet_cannot_expand_an_unbounded_rectangle() -> None:
    workbook = Workbook()
    assert workbook.active is not None
    workbook.active["XFD1048576"] = "tiny XML, huge logical sheet"
    with pytest.raises(ParseCorruptFile, match="bound"):
        await SheetParser().parse(_bytes_from(workbook.save), ParseContext())


async def test_bad_worksheet_dimensions_do_not_silently_drop_cells() -> None:
    package = _xlsx_fixture()
    with zipfile.ZipFile(io.BytesIO(package)) as archive:
        payload = archive.read("xl/worksheets/sheet1.xml").replace(b"A1:B2", b"A1:A1")
    parsed = await SheetParser().parse(
        _replace_member(package, "xl/worksheets/sheet1.xml", payload), ParseContext()
    )
    assert "=SUM(2,3)" in parsed.blocks[0].text


@pytest.mark.parametrize("text", ['name,value\nalpha,"unterminated', 'name,value\n"x"oops,1'])
async def test_malformed_csv_is_not_repaired_silently(text: str) -> None:
    with pytest.raises(ParseCorruptFile):
        await SheetParser().parse(text.encode(), ParseContext())


async def test_single_column_csv_is_valid() -> None:
    parsed = await SheetParser().parse(b"name\nalpha\n", ParseContext())
    assert parsed.markdown == "| name |\n| --- |\n| alpha |"


async def test_strict_mime_wrappers_never_switch_formats() -> None:
    with pytest.raises(ParseCorruptFile):
        await XlsxParser().parse(b"valid,csv\nnot,xlsx", ParseContext())
    parsed = await CsvParser().parse(b"PK,header\n1,2", ParseContext())
    assert "PK" in parsed.markdown
    assert "text/csv" not in XlsxParser.supported_mimes


def test_real_zip_encryption_flag_is_rejected() -> None:
    package = bytearray(_docx_fixture())
    central = package.index(b"PK\x01\x02")
    flags = struct.unpack_from("<H", package, central + 8)[0]
    struct.pack_into("<H", package, central + 8, flags | 1)
    with pytest.raises(ParseCorruptFile, match="Encrypted"):
        inspect_ooxml_package(bytes(package))


@pytest.mark.parametrize(
    "bound",
    ["MAX_INPUT_BYTES", "_MAX_MEMBERS", "_MAX_MEMBER_BYTES", "_MAX_TOTAL_BYTES", "_MAX_XML_NODES"],
)
def test_package_safety_bounds_are_enforced(bound: str, monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(f"cairn.ingestion.office_package.{bound}", 1)
    with pytest.raises(ParseCorruptFile):
        inspect_ooxml_package(_docx_fixture())


async def test_grouped_pptx_text_and_tables_are_not_dropped() -> None:
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    group = slide.shapes.add_group_shape()
    text_box = group.shapes.add_textbox(Inches(1), Inches(1), Inches(2), Inches(1))
    text_box.text = "Grouped content"
    parsed = await PptxParser().parse(_bytes_from(presentation.save), ParseContext())
    assert "Grouped content" in parsed.blocks[0].text


def test_docx_grid_span_cannot_expand_unbounded_cells() -> None:
    package = _docx_fixture()
    with zipfile.ZipFile(io.BytesIO(package)) as archive:
        payload = archive.read("word/document.xml").replace(
            b"<w:tcPr>", b'<w:tcPr><w:gridSpan w:val="999999999"/>', 1
        )
    with pytest.raises(ParseCorruptFile, match="bound"):
        inspect_ooxml_package(_replace_member(package, "word/document.xml", payload))


async def test_merged_word_cells_are_not_duplicated() -> None:
    document = Document()
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).merge(table.cell(0, 1)).text = "Merged heading"
    table.cell(1, 0).text = "value"
    parsed = await DocxParser().parse(_bytes_from(document.save), ParseContext())
    assert parsed.markdown.count("Merged heading") == 1


async def test_array_formula_source_is_not_an_object_representation() -> None:
    workbook = Workbook()
    assert workbook.active is not None
    workbook.active.append(["formula"])
    workbook.active["A2"] = ArrayFormula(ref="A2", text="=SUM(1,2)")
    parsed = await XlsxParser().parse(_bytes_from(workbook.save), ParseContext())
    assert "=SUM(1,2)" in parsed.markdown
    assert "object at" not in parsed.markdown


async def test_table_output_expansion_is_bounded(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr("cairn.ingestion.office._MAX_TABLE_CHARACTERS", 50, raising=False)
    with pytest.raises(ParseCorruptFile, match="bound"):
        await CsvParser().parse(b"header,value\n" + b"entry,123456789\n" * 10, ParseContext())
