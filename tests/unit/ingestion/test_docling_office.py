"""Word and PowerPoint through the shared Docling engine, with a recorded native fallback."""

from __future__ import annotations

import base64
import io
import json
import sys
import textwrap
from pathlib import Path
from uuid import uuid4

import pytest
from PIL import Image

from cairn.ingestion.base import ParseContext, ParsedDocument
from cairn.ingestion.docling_pdf import (
    DoclingDocxParser,
    DoclingPdfParser,
    DoclingPptxParser,
    normalize_docling,
)
from cairn.ingestion.ir import bind_document, digest, project_document
from cairn.ingestion.office import DocxParser, PptxParser
from cairn.ingestion.pdf import PdfEngineUnavailable, PdfLimits

DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


def _png() -> bytes:
    buffer = io.BytesIO()
    Image.new("RGB", (8, 8), "blue").save(buffer, format="PNG")
    return buffer.getvalue()


def _docx_bytes() -> bytes:
    from docx import Document

    document = Document()
    document.add_heading("Quarterly Report", level=1)
    document.add_paragraph("Revenue grew 12% in Q3.")
    table = document.add_table(rows=2, cols=2)
    table.rows[0].cells[0].text = "Region"
    table.rows[0].cells[1].text = "Revenue"
    table.rows[1].cells[0].text = "EMEA"
    table.rows[1].cells[1].text = "1200"
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _bind(parsed: ParsedDocument):  # type: ignore[no-untyped-def]  # reason: returns DocumentIR
    return bind_document(
        parsed,
        workspace_id=uuid4(),
        kb_id=uuid4(),
        document_id=uuid4(),
        revision=1,
        source_sha256="a" * 64,
        parser_fingerprint="b" * 64,
        context_fingerprint="c" * 64,
    )


def test_docx_payload_without_provenance_yields_unplaced_elements_and_a_picture_asset() -> None:
    png = _png()
    asset_id = __import__("hashlib").sha256(png).hexdigest()
    parsed = normalize_docling(
        {
            "format": "docx",
            "units": "points",
            "pages": {},
            "pdf_pages": [],
            "items": [
                {
                    "self_ref": "#/texts/0",
                    "label": "section_header",
                    "text": "Quarterly Report",
                    "prov": [],
                },
                {
                    "self_ref": "#/texts/1",
                    "label": "text",
                    "text": "Revenue grew 12% in Q3.",
                    "prov": [],
                },
                {
                    "self_ref": "#/pictures/0",
                    "label": "picture",
                    "text": "",
                    "prov": [],
                    "asset_id": asset_id,
                },
                {
                    "self_ref": "#/tables/0",
                    "label": "table",
                    "text": "",
                    "prov": [],
                    "data": {
                        "num_rows": 2,
                        "num_cols": 2,
                        "table_cells": [
                            {
                                "text": t,
                                "start_row_offset_idx": r,
                                "end_row_offset_idx": r + 1,
                                "start_col_offset_idx": c,
                                "end_col_offset_idx": c + 1,
                            }
                            for r, row in enumerate([["Region", "Revenue"], ["EMEA", "1200"]])
                            for c, t in enumerate(row)
                        ],
                    },
                },
            ],
            "assets": [
                {
                    "name": asset_id,
                    "mime_type": "image/png",
                    "data": base64.b64encode(png).decode(),
                    "page": None,
                }
            ],
        }
    )
    assert parsed.pages == 1
    assert parsed.metadata["format"] == "docx"
    assert "pdf_engine" not in parsed.metadata
    assert [element.type for element in parsed.elements] == ["heading", "text", "image", "table"]
    assert all(element.locations == () for element in parsed.elements)
    assert (
        parsed.elements[2].asset_id == asset_id and parsed.elements[2].asset_status == "available"
    )
    assert parsed.elements[3].table is not None and parsed.elements[3].table.markdown().startswith(
        "| Region | Revenue |"
    )
    assert parsed.elements[0].provenance[0].method == "native"
    ir = _bind(parsed)
    projected = project_document(ir)
    assert [block.page for block in projected.blocks] == [
        None,
        None,
        None,
    ]  # pictures carry no text
    assert ir.assets[0].mime_type == "image/png" and ir.assets[0].asset_id == asset_id


def test_pptx_emu_geometry_is_converted_to_points_bottom_left() -> None:
    parsed = normalize_docling(
        {
            "format": "pptx",
            "units": "emu",
            "pages": {"1": {"size": {"width": 9144000, "height": 6858000}}},
            "pdf_pages": [],
            "items": [
                {
                    "self_ref": "#/texts/0",
                    "label": "title",
                    "text": "Q3 results",
                    "prov": [
                        {
                            "page_no": 1,
                            "bbox": {
                                "l": 457200,
                                "b": 1417638,
                                "r": 8686800,
                                "t": 274638,
                                "coord_origin": "TOPLEFT",
                            },
                        }
                    ],
                }
            ],
        }
    )
    location = parsed.elements[0].locations[0]
    assert (location.width, location.height) == (720.0, 540.0)
    assert location.bbox == pytest.approx(
        (36.0, 540 - 1417638 / 12700, 684.0, 540 - 274638 / 12700)
    )
    assert location.bbox[1] < location.bbox[3]
    assert parsed.pages == 1 and parsed.metadata["format"] == "pptx"


class _FailingEngine:
    identity = "x" * 64

    async def convert(self, data: bytes, ctx: ParseContext, fmt: str) -> ParsedDocument:
        raise PdfEngineUnavailable()


class _RecordingEngine:
    identity = "y" * 64

    def __init__(self) -> None:
        self.calls: list[str] = []

    async def convert(self, data: bytes, ctx: ParseContext, fmt: str) -> ParsedDocument:
        self.calls.append(fmt)
        return ParsedDocument(
            markdown="from docling", blocks=[], metadata={"format": fmt, "engine": "docling"}
        )


async def test_office_parser_falls_back_to_the_native_parser_and_records_why() -> None:
    parser = DoclingDocxParser(engine=_FailingEngine(), fallback=DocxParser())
    parsed = await parser.parse(_docx_bytes(), ParseContext())
    assert "Quarterly Report" in parsed.markdown and "| Region | Revenue |" in parsed.markdown
    assert parsed.metadata["docling_fallback"] == {
        "reason": "PDF_ENGINE_UNAVAILABLE",
        "parser": "DocxParser",
    }
    assert any("Docling conversion failed" in w for w in parsed.metadata["warnings"])
    assert parser.fallback_implementation == "cairn/DocxParser/3"


async def test_office_parser_uses_docling_and_passes_the_format() -> None:
    engine = _RecordingEngine()
    docx = DoclingDocxParser(engine=engine, fallback=DocxParser())
    pptx = DoclingPptxParser(engine=engine, fallback=PptxParser())
    parsed = await docx.parse(_docx_bytes(), ParseContext())
    assert parsed.markdown == "from docling" and "docling_fallback" not in parsed.metadata
    assert engine.calls == ["docx"]
    assert pptx.format == "pptx" and DOCX_MIME in docx.supported_mimes


FAKE_CHILD = textwrap.dedent(
    """
    import json, os, struct, sys

    config = json.loads(sys.argv[1])
    inp, out = sys.stdin.buffer, sys.stdout.buffer

    def read_frame():
        header = inp.read(4)
        if len(header) < 4:
            return None
        (length,) = struct.unpack(">I", header)
        data = b""
        while len(data) < length:
            chunk = inp.read(length - len(data))
            if not chunk:
                return None
            data += chunk
        return data

    def write_frame(data):
        out.write(struct.pack(">I", len(data)) + data)
        out.flush()

    write_frame(json.dumps({"ready": True, "identity": config["identity"]}).encode())
    while True:
        meta = read_frame()
        if meta is None:
            break
        source = read_frame()
        if source is None:
            break
        request = json.loads(meta)
        fmt = request.get("format", "pdf")
        payload = {
            "request_id": request["request_id"],
            "identity": config["identity"],
            "worker_pid": os.getpid(),
            "format": fmt,
            "units": "points",
            "pages": {},
            "pdf_pages": [],
            "items": [
                {"self_ref": "#/texts/0", "label": "text", "text": "hello " + fmt, "prov": []}
            ],
            "pdf_budget_facts": {"schema": 1, "pages": [], "output_bytes": 100},
        }
        write_frame(json.dumps(payload).encode())
    """
)


async def test_persistent_engine_sends_the_format_and_skips_pdf_checks_for_office(
    tmp_path: Path,
) -> None:
    manifest = {"docling": "2.135.0", "files": [], "packages": {}}
    manifest["identity"] = digest(manifest)
    (tmp_path / "manifest.json").write_text(json.dumps(manifest))
    (tmp_path / "fake.py").write_text(FAKE_CHILD, encoding="utf-8")
    engine = DoclingPdfParser(
        python=sys.executable,
        assets_path=str(tmp_path),
        manifest_path=str(tmp_path / "manifest.json"),
        ocr=False,
        languages=("eng",),
        persistent=True,
        limits=PdfLimits(timeout_seconds=5, max_output_bytes=50_000, max_memory_bytes=4 * 1024**3),
    )
    engine._worker_script = str(tmp_path / "fake.py")
    try:
        parsed = await engine.convert(b"PK\x03\x04 not a pdf", ParseContext(), "docx")
    finally:
        await engine.aclose()
    assert parsed.metadata["format"] == "docx"
    assert parsed.markdown == "hello docx"
    assert parsed.elements[0].locations == ()
