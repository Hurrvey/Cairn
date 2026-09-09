"""Office fixtures through real queue, catalog, object storage and pgvector."""

from __future__ import annotations

from io import BytesIO
from pathlib import Path

import pytest
from docx import Document
from openpyxl import Workbook
from pptx import Presentation
from sqlalchemy import text
from tests.integration.test_ingestion_pipeline import (
    _create_pipeline_document,
    _execute_pipeline,
    _test_pipeline,
    pipeline_admin,
)

from cairn.authz.model import Principal
from cairn.catalog.service import CatalogService
from cairn.core.db import session_scope, transaction
from cairn.ingestion.artifacts import decode_parse_manifest
from cairn.vectorstore.base import Namespace

__all__ = ["pipeline_admin"]


def _office_source(extension: str) -> tuple[bytes, str, str]:
    buffer = BytesIO()
    if extension == "docx":
        document = Document()
        document.add_paragraph(
            "Document processing preserves the original source and page references."
        )
        document.save(buffer)
        return (
            buffer.getvalue(),
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            "source and page references",
        )
    if extension == "pptx":
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[1])
        assert slide.shapes.title is not None
        slide.shapes.title.text = "Report"
        slide.placeholders[1].text = "Original source references"
        presentation.save(buffer)
        return (
            buffer.getvalue(),
            "application/vnd.openxmlformats-officedocument.presentationml.presentation",
            "Original source references",
        )
    if extension == "xlsx":
        workbook = Workbook()
        worksheet = workbook.active
        assert worksheet is not None
        worksheet.append(["Item", "Total"])
        worksheet.append(["Books", 7])
        workbook.save(buffer)
        workbook.close()
        return (
            buffer.getvalue(),
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            "Books",
        )
    return b"Item,Total\nBooks,7\n", "text/csv", "Books"


@pytest.mark.parametrize("extension", ["docx", "pptx", "xlsx", "csv"])
async def test_office_upload_reaches_verified_index(
    pipeline_admin: Principal, tmp_path: Path, extension: str
) -> None:
    source, mime, expected = _office_source(extension)
    catalog, kb, registration, _source, _key = await _create_pipeline_document(
        pipeline_admin, tmp_path, source=source
    )
    assert registration.document is not None
    async with transaction() as session:
        await session.execute(
            text("UPDATE document SET mime_type=:mime WHERE id=:document_id"),
            {"mime": mime, "document_id": registration.document.id},
        )
    (
        pipeline,
        provider,
        objects,
        vectors,
        _resolved_objects,
        _resolved_vectors,
    ) = await _test_pipeline(kb, tmp_path)
    await _execute_pipeline(pipeline)
    document = await catalog.get_document(registration.document.id)
    assert document.state == "indexed"
    assert (await CatalogService().get_kb(kb.id)).active_index_version == 1
    assert await vectors.count(Namespace(kb.id, 1)) > 0
    assert any(expected in content for batch in provider.calls for content in batch)
    async with session_scope() as session:
        parsed_key = await session.scalar(
            text("SELECT parsed_object_key FROM document_ingestion WHERE document_id=:document_id"),
            {"document_id": document.id},
        )
    assert isinstance(parsed_key, str)
    parsed = decode_parse_manifest(await objects.get_bytes(parsed_key, max_bytes=1024 * 1024))
    assert expected in parsed.parsed.markdown
    if extension == "docx":
        assert parsed.parsed.language == "en"
        assert all(block.page is None for block in parsed.parsed.blocks)
    else:
        assert all(block.page == 1 for block in parsed.parsed.blocks)
