"""Typed boundaries for the OpenPyXL reader and defused XML event stream."""

from __future__ import annotations

import io
from collections.abc import Iterator
from typing import Protocol, cast
from xml.etree.ElementTree import Element

from defusedxml.ElementTree import iterparse
from openpyxl import load_workbook
from openpyxl.worksheet.formula import ArrayFormula, DataTableFormula

from cairn.ingestion.errors import ParseCorruptFile


class ReadOnlyWorksheet(Protocol):
    title: str
    sheet_state: str

    def reset_dimensions(self) -> None: ...
    def iter_rows(self, *, values_only: bool) -> Iterator[tuple[object, ...]]: ...


class ReadOnlyWorkbook(Protocol):
    @property
    def worksheets(self) -> list[ReadOnlyWorksheet]: ...

    def close(self) -> None: ...


def load_read_only_workbook(data: bytes) -> ReadOnlyWorkbook:
    workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=False, keep_links=False)
    return cast("ReadOnlyWorkbook", workbook)


def safe_xml_events(data: bytes) -> Iterator[tuple[str, Element]]:
    events = iterparse(
        io.BytesIO(data),
        events=("start", "end"),
        forbid_dtd=True,
        forbid_entities=True,
        forbid_external=True,
    )
    return cast("Iterator[tuple[str, Element]]", events)


def formula_source(value: object) -> object:
    if isinstance(value, DataTableFormula):
        raise ParseCorruptFile("Excel data-table formulas do not contain indexable formula source.")
    if isinstance(value, ArrayFormula):
        if not isinstance(value.text, str) or not value.text:
            raise ParseCorruptFile("An Excel array formula has no source expression.")
        return value.text
    return value
