"""Offline DOCX, PPTX, XLSX, and CSV parser adapters."""

from __future__ import annotations

import csv
import io
from collections.abc import Iterable, Iterator
from datetime import date, datetime, time
from typing import ClassVar

from anyio import to_thread
from docx import Document
from docx.document import Document as DocumentType
from docx.table import Table as DocxTable
from docx.text.paragraph import Paragraph
from pptx import Presentation
from pptx.presentation import Presentation as PresentationType
from pptx.shapes.autoshape import Shape
from pptx.shapes.base import BaseShape
from pptx.shapes.graphfrm import GraphicFrame
from pptx.shapes.group import GroupShape
from pydantic import JsonValue

from cairn.ingestion.base import Block, BlockKind, ParseContext, ParsedDocument
from cairn.ingestion.errors import ParseCorruptFile, ParseEmptyContent, ParseError
from cairn.ingestion.office_compat import ReadOnlyWorksheet, formula_source, load_read_only_workbook
from cairn.ingestion.office_package import (
    MAX_INPUT_BYTES,
    MAX_TABLE_CELLS,
    MAX_TABLE_COLUMNS,
    MAX_TABLE_ROWS,
    inspect_ooxml_package,
)

_DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
_PPTX_MIME = "application/vnd.openxmlformats-officedocument.presentationml.presentation"
_XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
_MAX_TABLE_CHARACTERS = 8 * 1024 * 1024
_MAX_DOCUMENT_CHARACTERS = 16 * 1024 * 1024


class DocxParser:
    supported_mimes: ClassVar[frozenset[str]] = frozenset({_DOCX_MIME})

    async def parse(self, data: bytes, ctx: ParseContext) -> ParsedDocument:
        return await to_thread.run_sync(self._parse, data, ctx)

    def _parse(self, data: bytes, ctx: ParseContext) -> ParsedDocument:
        package = inspect_ooxml_package(data)
        try:
            document = Document(io.BytesIO(package.data))
            blocks = _docx_blocks(document, ctx)
        except ParseError:
            raise
        except Exception as exc:
            raise ParseCorruptFile() from exc
        return _document_from_blocks(blocks, ctx, pages=1, metadata={"format": "docx"})


class PptxParser:
    supported_mimes: ClassVar[frozenset[str]] = frozenset({_PPTX_MIME})

    async def parse(self, data: bytes, ctx: ParseContext) -> ParsedDocument:
        return await to_thread.run_sync(self._parse, data, ctx)

    def _parse(self, data: bytes, ctx: ParseContext) -> ParsedDocument:
        package = inspect_ooxml_package(data)
        try:
            presentation = Presentation(io.BytesIO(package.data))
            blocks, slides = _pptx_blocks(presentation, ctx)
        except ParseError:
            raise
        except Exception as exc:
            raise ParseCorruptFile() from exc
        return _document_from_blocks(
            blocks,
            ctx,
            pages=len(presentation.slides),
            metadata={"format": "pptx", "slides": slides},
        )


class SheetParser:
    supported_mimes: ClassVar[frozenset[str]] = frozenset({_XLSX_MIME, "text/csv"})

    async def parse(self, data: bytes, ctx: ParseContext) -> ParsedDocument:
        return await to_thread.run_sync(self._parse, data, ctx)

    def _parse(self, data: bytes, ctx: ParseContext) -> ParsedDocument:
        if data.startswith((b"PK", b"\xd0\xcf\x11\xe0")):
            return self._parse_xlsx(data, ctx)
        return self._parse_csv(data, ctx)

    def _parse_xlsx(self, data: bytes, ctx: ParseContext) -> ParsedDocument:
        package = inspect_ooxml_package(data)
        workbook = None
        try:
            workbook = load_read_only_workbook(package.data)
            blocks: list[Block] = []
            sheet_metadata: list[JsonValue] = []
            for sheet_number, worksheet in enumerate(workbook.worksheets, start=1):
                sheet_metadata.append({"name": worksheet.title, "state": worksheet.sheet_state})
                rows = _worksheet_rows(worksheet)
                if not rows:
                    continue
                table = _table_markdown(rows)
                blocks.append(
                    Block(
                        kind="table",
                        text=table,
                        ordinal=len(blocks),
                        heading_path=(worksheet.title,),
                        heading_level=1,
                        page=sheet_number,
                        source_url=ctx.source_url,
                    )
                )
            return _document_from_blocks(
                blocks,
                ctx,
                pages=len(workbook.worksheets),
                metadata={
                    "format": "xlsx",
                    "formula_mode": "source",
                    "sheets": sheet_metadata,
                },
            )
        except ParseError:
            raise
        except Exception as exc:
            raise ParseCorruptFile() from exc
        finally:
            if workbook is not None:
                workbook.close()

    def _parse_csv(self, data: bytes, ctx: ParseContext) -> ParsedDocument:
        if len(data) > MAX_INPUT_BYTES:
            raise ParseCorruptFile("The CSV exceeds the parser safety bound.")
        try:
            text = data.decode(ctx.encoding).removeprefix("\ufeff")
        except (LookupError, UnicodeError) as exc:
            raise ParseCorruptFile() from exc
        if "\x00" in text:
            raise ParseCorruptFile()
        if not text.strip():
            raise ParseEmptyContent()
        delimiter = _csv_delimiter(text)
        try:
            rows: list[list[str]] = []
            width = 0
            for row in csv.reader(io.StringIO(text, newline=""), delimiter=delimiter, strict=True):
                if not row:
                    continue
                width = max(width, len(row))
                if (
                    width > MAX_TABLE_COLUMNS
                    or len(rows) >= MAX_TABLE_ROWS
                    or (len(rows) + 1) * width > MAX_TABLE_CELLS
                ):
                    raise ParseCorruptFile("The CSV exceeds the table safety bound.")
                rows.append(row)
        except (csv.Error, UnicodeError) as exc:
            raise ParseCorruptFile() from exc
        if not any(cell.strip() for row in rows for cell in row):
            raise ParseEmptyContent()
        table = _table_markdown(rows)
        block = Block(
            kind="table",
            text=table,
            ordinal=0,
            heading_level=1,
            page=1,
            source_url=ctx.source_url,
        )
        return ParsedDocument(
            markdown=table,
            blocks=[block],
            pages=1,
            language=ctx.language,
            metadata={"format": "csv", "delimiter": delimiter, "encoding": ctx.encoding},
        )


class XlsxParser(SheetParser):
    """MIME-bound XLSX adapter; corrupt packages cannot fall back to CSV."""

    supported_mimes: ClassVar[frozenset[str]] = frozenset({_XLSX_MIME})

    def _parse(self, data: bytes, ctx: ParseContext) -> ParsedDocument:
        return self._parse_xlsx(data, ctx)


class CsvParser(SheetParser):
    """MIME-bound CSV adapter; a ZIP signature is not a format-switch command."""

    supported_mimes: ClassVar[frozenset[str]] = frozenset({"text/csv"})

    def _parse(self, data: bytes, ctx: ParseContext) -> ParsedDocument:
        return self._parse_csv(data, ctx)


def _csv_delimiter(text: str) -> str:
    sample = text[: 64 * 1024]
    try:
        return csv.Sniffer().sniff(sample, delimiters=",;\t").delimiter
    except csv.Error:
        first_line = sample.splitlines()[0]
        try:
            return csv.Sniffer().sniff(first_line, delimiters=",;\t").delimiter
        except csv.Error:
            return ","


def _docx_blocks(document: DocumentType, ctx: ParseContext) -> list[Block]:
    blocks: list[Block] = []
    headings: list[tuple[int, str]] = []
    for item in document.iter_inner_content():
        if isinstance(item, Paragraph):
            text = item.text.strip()
            if not text:
                continue
            style_name = item.style.name if item.style is not None else ""
            level = _heading_level(style_name)
            if level is not None:
                while headings and headings[-1][0] >= level:
                    headings.pop()
                headings.append((level, text))
                kind: BlockKind = "heading"
                rendered = f"{'#' * min(level, 6)} {text}"
            elif style_name.lower().startswith("list "):
                kind = "list"
                marker = "1." if style_name.lower().startswith("list number") else "-"
                rendered = f"{marker} {text}"
            else:
                kind = "paragraph"
                rendered = text
            blocks.append(
                Block(
                    kind=kind,
                    text=rendered,
                    ordinal=len(blocks),
                    heading_path=tuple(title for _, title in headings),
                    heading_level=level,
                    source_url=ctx.source_url,
                )
            )
            continue
        if isinstance(item, DocxTable):
            rows = _docx_table_rows(item)
            if any(cell for row in rows for cell in row):
                blocks.append(
                    Block(
                        kind="table",
                        text=_table_markdown(rows),
                        ordinal=len(blocks),
                        heading_path=tuple(title for _, title in headings),
                        page=None,
                        source_url=ctx.source_url,
                    )
                )
    return blocks


def _docx_table_rows(table: DocxTable) -> list[list[str]]:
    seen_cells: set[object] = set()
    rows: list[list[str]] = []
    for row in table.rows:
        values: list[str] = [""] * row.grid_cols_before
        for cell in row.cells:
            values.append("" if cell._tc in seen_cells else cell.text.strip())
            seen_cells.add(cell._tc)
        values.extend([""] * row.grid_cols_after)
        rows.append(values)
    return rows


def _leaf_shapes(shapes: Iterable[BaseShape]) -> Iterator[BaseShape]:
    for shape in shapes:
        if isinstance(shape, GroupShape):
            yield from _leaf_shapes(shape.shapes)
        else:
            yield shape


def _pptx_blocks(
    presentation: PresentationType, ctx: ParseContext
) -> tuple[list[Block], list[JsonValue]]:
    blocks: list[Block] = []
    slides: list[JsonValue] = []
    for slide_number, slide in enumerate(presentation.slides, start=1):
        title_shape = slide.shapes.title
        title = title_shape.text.strip() if title_shape is not None else ""
        slides.append({"number": slide_number, "title": title})
        parts: list[str] = []
        has_table = False
        for shape in _leaf_shapes(slide.shapes):
            if title_shape is not None and shape == title_shape:
                continue
            if isinstance(shape, GraphicFrame) and shape.has_table:
                has_table = True
                table_rows = [[cell.text.strip() for cell in row.cells] for row in shape.table.rows]
                if any(cell for row in table_rows for cell in row):
                    parts.append(_table_markdown(table_rows))
            elif isinstance(shape, Shape) and shape.has_text_frame:
                text = shape.text.strip()
                if text:
                    parts.append(text)
        if title:
            parts.insert(0, f"# {title}")
        text = "\n\n".join(parts).strip()
        if not text:
            continue
        blocks.append(
            Block(
                kind="table" if has_table else "paragraph",
                text=text,
                ordinal=len(blocks),
                heading_path=(title,) if title else (),
                heading_level=1 if title else None,
                page=slide_number,
                source_url=ctx.source_url,
            )
        )
    return blocks, slides


def _worksheet_rows(worksheet: ReadOnlyWorksheet) -> list[list[object]]:
    worksheet.reset_dimensions()
    rows: list[list[object]] = []
    for row_number, values in enumerate(worksheet.iter_rows(values_only=True), start=1):
        if row_number > MAX_TABLE_ROWS:
            raise ParseCorruptFile("The worksheet exceeds the row safety bound.")
        row = [formula_source(value) for value in values]
        if len(row) > MAX_TABLE_COLUMNS:
            raise ParseCorruptFile("The worksheet exceeds the column safety bound.")
        rows.append(row)
    while rows and not any(value not in (None, "") for value in rows[-1]):
        rows.pop()
    if not rows or not any(value not in (None, "") for row in rows for value in row):
        return []
    width = max(len(row) for row in rows)
    if width * len(rows) > MAX_TABLE_CELLS:
        raise ParseCorruptFile("The worksheet exceeds the cell safety bound.")
    return [row + [None] * (width - len(row)) for row in rows]


def _heading_level(style_name: str) -> int | None:
    prefix, _, suffix = style_name.partition(" ")
    if prefix.lower() != "heading" or not suffix.isdigit():
        return None
    level = int(suffix)
    return level if 1 <= level <= 9 else None


def _table_markdown(rows: Iterable[Iterable[object]]) -> str:
    normalized = [list(row) for row in rows]
    if not normalized:
        raise ParseEmptyContent()
    width = max(len(row) for row in normalized)
    if width > MAX_TABLE_COLUMNS or width * len(normalized) > MAX_TABLE_CELLS:
        raise ParseCorruptFile("The table exceeds the cell safety bound.")
    rectangular = [row + [None] * (width - len(row)) for row in normalized]
    lines: list[str] = []
    characters = 0
    for row_number, row in enumerate(rectangular):
        cells: list[str] = []
        for value in row:
            text = _cell_text(value)
            characters += len(text) + 3
            if characters > _MAX_TABLE_CHARACTERS:
                raise ParseCorruptFile("The rendered table exceeds the output safety bound.")
            cells.append(text)
        lines.append("| " + " | ".join(cells) + " |")
        if row_number == 0:
            separator = "| " + " | ".join("---" for _ in range(width)) + " |"
            characters += len(separator) + 1
            lines.append(separator)
    return "\n".join(lines)


def _cell_text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        text = value.isoformat(sep=" ")
    elif isinstance(value, (date, time)):
        text = value.isoformat()
    else:
        text = str(value)
    if len(text) > _MAX_TABLE_CHARACTERS:
        raise ParseCorruptFile("The cell exceeds the output safety bound.")
    return (
        text.replace("&", "&amp;")
        .replace("<", "&lt;")
        .replace(">", "&gt;")
        .replace("\\", "\\\\")
        .replace("|", "\\|")
        .replace("\r\n", "\n")
        .replace("\r", "\n")
        .replace("\n", "<br>")
    )


def _document_from_blocks(
    blocks: list[Block],
    ctx: ParseContext,
    *,
    pages: int,
    metadata: dict[str, JsonValue],
) -> ParsedDocument:
    if not blocks:
        raise ParseEmptyContent()
    if sum(len(block.text) + 2 for block in blocks) > _MAX_DOCUMENT_CHARACTERS:
        raise ParseCorruptFile("The document exceeds the output safety bound.")
    markdown = "\n\n".join(block.text for block in blocks)
    return ParsedDocument(
        markdown=markdown,
        blocks=blocks,
        pages=max(1, pages),
        language=ctx.language,
        metadata=metadata,
    )
